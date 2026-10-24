"""BLS-Excel-Datei → schlanke CSV für den Import ins Add-on.

Läuft am Rechner, nicht im Add-on:

    pip install -r tools/requirements.txt
    python tools/bls_convert.py dev/bls/BLS_4_0_2025_DE/BLS_4_0_Daten_2025_DE.xlsx \\
        tellerplan/app/data/bls_4_0.csv

Übernommen werden BLS-Code, Name und sieben Nährwerte pro 100 g. Angaben wie
„TR“ (Spuren) oder „<LOD“ (unter der Nachweisgrenze) werden zu 0, „-“ (unbekannt)
bleibt leer.

Quelle: Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), Version 4.0 –
Deutsche Nährstoffdatenbank. Lizenz CC BY 4.0.
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

# CSV-Spalte → BLS-Nährstoffcode
NUTRIENT_CODES = {
    "kcal": "ENERCC",
    "protein_g": "PROT625",
    "fat_g": "FAT",
    "carbs_g": "CHO",
    "fiber_g": "FIBT",
    "sugar_g": "SUGAR",
    "salt_g": "NACL",
}
CODE_COLUMN = "BLS Code"
NAME_COLUMN = "Lebensmittelbezeichnung"
BELOW_DETECTION = {"TR", "<LOD", "<LOQ", "<LOD or <LOQ"}
UNKNOWN = {"-", ""}


def _value_columns(header: tuple) -> dict[str, int]:
    """Spaltenindex je Nährstoffcode (nur die Wertspalte, nicht Herkunft/Referenz)."""
    columns = {}
    for index, title in enumerate(header):
        text = str(title)
        if "Datenherkunft" in text or "Referenz" in text:
            continue
        columns[text.split(" ")[0]] = index
    return columns


def _number(value: object) -> str:
    if value is None or (isinstance(value, str) and value.strip() in UNKNOWN):
        return ""
    if isinstance(value, str) and value.strip() in BELOW_DETECTION:
        return "0"
    if isinstance(value, int | float):
        return f"{value:g}"
    raise ValueError(f"Unerwarteter Wert: {value!r}")


def convert(source: Path, target: Path) -> int:
    sheet = load_workbook(source, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = _value_columns(header)
    code_index = header.index(CODE_COLUMN)
    name_index = header.index(NAME_COLUMN)

    records = []
    for row in rows:
        if not row[code_index]:
            continue
        record = {"bls_code": row[code_index], "name": row[name_index].strip()}
        for field, code in NUTRIENT_CODES.items():
            record[field] = _number(row[columns[code]])
        records.append(record)
    records.sort(key=lambda record: record["bls_code"])

    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["bls_code", "name", *NUTRIENT_CODES])
        writer.writeheader()
        writer.writerows(records)
    return len(records)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    count = convert(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"{count} Lebensmittel nach {sys.argv[2]} geschrieben.")
